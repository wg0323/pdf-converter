import os
import gc
import sys
import time
import ctypes
import logging
import queue
import traceback
import tempfile
import multiprocessing
from pathlib import Path
from typing import Optional, TYPE_CHECKING

# pdf2docx 导入链很重（PyMuPDF/numpy/opencv），懒加载以免拖慢主程序启动
if TYPE_CHECKING:
    from pdf2docx import Converter

logger = logging.getLogger(__name__)

# 扫描版PDF分卷转换的每卷页数：一次性把整本（如906页/200万字）构建进
# python-docx 会撑爆内存（lxml 文档树 + 保存序列化放大数倍），故每 50 页
# 独立生成一个 docx 卷并立即落盘、释放，把内存峰值限制在单卷规模。
_OCR_DOCX_CHUNK_PAGES = 50

# OCR 子进程 → 父进程 的消息类型与结果哨兵
_OCR_MSG_PROGRESS = 'progress'      # (页进度: 当前页, 总页数)
_OCR_MSG_LOG = 'log'                # (用户可见日志文本)
_OCR_MSG_RESULT = 'result'          # (最终结果: 是否成功, 消息)
# OCR 不可用（增强包缺失/依赖损坏/引擎初始化失败），父进程据此回退图片模式
_OCR_UNAVAILABLE = '__OCR_UNAVAILABLE__'
# 子进程原生崩溃（exitcode 非0），父进程据此尝试关闭 MKLDNN 重跑
_OCR_CRASH = '__OCR_CRASH__'


def _pdf2docx_worker(pdf_path, output_path, start, end, result_queue):
    """
    在独立子进程中执行 pdf2docx 转换。

    pdf2docx 依赖的 PyMuPDF C 扩展在解析某些文档时可能触发 refcount 错误
    等原生崩溃，这类崩溃无法被 Python 的 try/except 捕获，会直接终止进程。
    放到子进程执行后，即使崩溃也只是子进程退出，父进程（GUI）通过 exitcode
    感知并优雅报错，不会被整体拖垮。
    """
    try:
        from pdf2docx import Converter
        cv = Converter(pdf_path)
        cv.convert(
            output_path,
            start=start,
            end=end,
            pages=None,
            debug=False,
            multi_processing=False,
        )
        cv.close()
        result_queue.put((True, ""))
    except Exception as e:
        result_queue.put((False, f"转换失败: {e}"))


def _ocr_cpu_threads() -> int:
    """OCR CPU 推理线程数：对齐物理核估计（逻辑核数的一半，至少 2）。

    超线程对 MKLDNN 数值计算无收益，线程数超过物理核反而增加调度开销。
    """
    return max(2, (os.cpu_count() or 8) // 2)


def _find_ocr_addon_dir() -> Optional[str]:
    """
    查找 OCR 增强包目录（ocr_addon）

    增强包内含 paddle/paddleocr 及其依赖与离线模型，单独分发，
    用户解压到程序目录内即可启用 OCR。
    打包环境找 exe 同级目录，开发环境找项目根目录。

    Returns:
        增强包目录路径，不存在则返回 None
    """
    if getattr(sys, 'frozen', False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = str(Path(__file__).resolve().parents[2])
    addon_dir = os.path.join(base_dir, 'ocr_addon')
    return addon_dir if os.path.isdir(addon_dir) else None


def _setup_ocr_addon_paths(addon_dir: str):
    """
    把增强包 site-packages 注入 sys.path 并修正 site.USER_SITE。

    依赖放在 site-packages 子目录：paddle 启动时按路径名包含
    "site-packages" 定位原生库目录，否则在冻结环境回退到
    site.USER_SITE（为 None）而崩溃；PyInstaller 冻结环境下
    site.USER_SITE 为 None，而 paddle 启动时会用它拼接原生库路径
    （不判空，直接 TypeError），指向增强包 site-packages 使其能
    定位 paddle/libs。
    """
    addon_site = os.path.join(addon_dir, 'site-packages')
    if os.path.isdir(addon_site) and addon_site not in sys.path:
        sys.path.insert(0, addon_site)
    import site
    if getattr(site, 'USER_SITE', None) is None:
        site.USER_SITE = addon_site


def _create_ocr_engine(addon_dir: str, use_mkldnn: bool):
    """
    创建 PP-Structure 引擎（仅在 OCR 子进程内调用）

    PP-Structure = 版面分析 + OCR + 阅读顺序恢复，可直出分层可编辑的 Word。
    优先使用增强包内的离线模型，避免首次使用时联网下载；
    开启 MKLDNN 加速 CPU 推理，线程数对齐物理核。
    """
    _setup_ocr_addon_paths(addon_dir)
    from paddleocr import PPStructure

    model_kwargs = {}
    models_dir = os.path.join(addon_dir, 'models')
    for name, param in (
        ('det', 'det_model_dir'),
        ('rec', 'rec_model_dir'),
        ('layout', 'layout_model_dir'),
    ):
        model_dir = os.path.join(models_dir, name)
        if os.path.isdir(model_dir):
            model_kwargs[param] = model_dir

    # table=False：表格模型与 paddlepaddle 2.6 存在加载兼容问题，
    # 表格区域按图片处理；recovery=True 启用阅读顺序恢复；
    # rec_batch_num=24：密集页一页 det 出 40+ 文本行，默认批 6 要跑 7 批
    # （调度开销占比高），提到 24 后实测页均 2.07s→1.67s（-19%，纯批调度
    # 优化，识别结果不变）
    return PPStructure(
        recovery=True,
        table=False,
        lang='ch',
        show_log=False,
        enable_mkldnn=use_mkldnn,
        cpu_threads=_ocr_cpu_threads(),
        rec_batch_num=24,
        **model_kwargs
    )


def _ocr_worker(pdf_path, output_path, start, end, msg_queue, use_mkldnn):
    """
    在独立子进程中执行扫描版 PDF 的 PP-Structure OCR 转换。

    - 低优先级运行（Windows BELOW_NORMAL）：paddle CPU 推理会吃满所有
      核心，降低进程优先级保证转换期间 GUI 与其他前台程序仍然可用；
    - 通过 msg_queue 向父进程发送 progress/log/result 三类消息；
    - 取消由父进程 terminate 子进程实现，无需协作式检查。

    转换主体：每页渲染为图片 → 版面分析（标题/正文/图片等区域）→ OCR
    识别 → 阅读顺序恢复 → 官方 recovery 直出分层可编辑的 docx。
    为避免超大文档（如906页）一次性构建 docx 撑爆内存，按每
    _OCR_DOCX_CHUNK_PAGES 页分卷生成：第一卷用原始输出名，第2卷起追加
    _partNN 后缀，各卷独立落盘并释放内存。
    """
    # Windows：降低进程优先级，OCR 满载时前台程序不受拖累
    if sys.platform == 'win32':
        try:
            kernel32 = ctypes.windll.kernel32
            # 0x4000 = BELOW_NORMAL_PRIORITY_CLASS
            kernel32.SetPriorityClass(kernel32.GetCurrentProcess(), 0x4000)
        except Exception:
            pass

    def log(message):
        msg_queue.put((_OCR_MSG_LOG, message))

    def finish(success, message):
        msg_queue.put((_OCR_MSG_RESULT, success, message))

    addon_dir = _find_ocr_addon_dir()
    if addon_dir is None:
        finish(False, _OCR_UNAVAILABLE)
        return

    try:
        t_init = time.time()
        engine = _create_ocr_engine(addon_dir, use_mkldnn)
        log(
            f"已启用 OCR 文字识别（PP-Structure，MKLDNN {'开' if use_mkldnn else '关'}，"
            f"CPU 线程数 {_ocr_cpu_threads()}），开始逐页识别..."
            f"（引擎初始化 {time.time() - t_init:.1f}s）"
        )
    except ImportError as e:
        # 半初始化的 paddle 也可能以 ImportError 形式抛出（增强包不完整）
        log(f"OCR 增强包不完整（缺少 paddle 依赖：{e}）")
        finish(False, _OCR_UNAVAILABLE)
        return
    except Exception as e:
        log(f"OCR 引擎初始化失败：{e}")
        finish(False, _OCR_UNAVAILABLE)
        return

    import fitz
    import shutil
    import numpy as np
    import cv2
    from paddleocr import save_structure_res
    from paddleocr.ppstructure.recovery.recovery_to_doc import (
        sorted_layout_boxes, convert_info_docx,
    )

    doc = None
    temp_dir = None
    doc_name = 'ppstructure_result'
    zoom = 200 / 72  # 200 DPI
    t_total = time.time()
    infer_seconds = 0.0
    pages_done = 0

    try:
        doc = fitz.open(pdf_path)
        total_pages = len(doc)
        actual_end = min(end if end is not None else total_pages, total_pages)
        pages_to_convert = actual_end - start

        logger.info(f"扫描版PDF，使用PP-Structure转换: 共{total_pages}页，转换第{start + 1}-{actual_end}页")

        # 分卷转换：每 _OCR_DOCX_CHUNK_PAGES 页独立构建一个 docx 卷并立即
        # 落盘、释放该卷内存。避免把整本（如906页/200万字）一次性塞进
        # python-docx —— lxml 文档树加保存序列化会把内存放大数倍而撑爆。
        chunk_starts = list(range(start, actual_end, _OCR_DOCX_CHUNK_PAGES))
        multi_volume = len(chunk_starts) > 1
        out_stem = os.path.splitext(output_path)[0]
        if multi_volume:
            log(f"文档较大（{pages_to_convert}页），按每{_OCR_DOCX_CHUNK_PAGES}页分卷输出，共{len(chunk_starts)}卷")

        output_files = []
        for vol_idx, chunk_start in enumerate(chunk_starts, start=1):
            chunk_end = min(chunk_start + _OCR_DOCX_CHUNK_PAGES, actual_end)

            # 每卷独立 ASCII 临时目录：save_structure_res 的 img_idx 卷内从0起，
            # 独立目录避免跨卷图名冲突，卷末即清理释放磁盘与内存
            # （cv2.imwrite 在 Windows 无法写含中文路径，故必须用 ASCII 临时目录）
            temp_dir = tempfile.mkdtemp(prefix='pdf_ppstructure_')
            chunk_res = []
            last_img = None
            try:
                for page_idx in range(chunk_start, chunk_end):
                    current_page = page_idx - start + 1
                    msg_queue.put((_OCR_MSG_PROGRESS, current_page, pages_to_convert))

                    logger.info(f"PP-Structure识别第 {page_idx + 1}/{total_pages} 页...")

                    # 渲染页面为 BGR ndarray
                    pix = doc[page_idx].get_pixmap(matrix=fitz.Matrix(zoom, zoom))
                    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
                        pix.height, pix.width, pix.n
                    )
                    if pix.n == 4:
                        img = cv2.cvtColor(img, cv2.COLOR_RGBA2BGR)
                    else:
                        img = cv2.cvtColor(img, cv2.COLOR_RGB2BGR)

                    # 版面分析 + OCR；rel_idx 卷内从0，用于图名与 docx 分页
                    rel_idx = page_idx - chunk_start
                    t_page = time.time()
                    result = engine(img, img_idx=rel_idx)
                    infer_seconds += time.time() - t_page
                    pages_done += 1

                    # 图片区域裁剪落盘，convert_info_docx 生成 docx 时需要读取
                    save_structure_res(result, temp_dir, doc_name, img_idx=rel_idx)
                    res = sorted_layout_boxes(result, img.shape[1])
                    chunk_res += res
                    last_img = img

                if not chunk_res:
                    logger.info(f"第{vol_idx}卷（第{chunk_start + 1}-{chunk_end}页）无识别内容，跳过")
                    continue

                # 仅构建本卷 docx（内存峰值≈单卷），随即落盘
                convert_info_docx(last_img, chunk_res, temp_dir, doc_name)
                temp_docx = os.path.join(temp_dir, f'{doc_name}_ocr.docx')
                if not os.path.exists(temp_docx):
                    finish(False, f"第{vol_idx}卷转换完成但输出文件不存在")
                    return

                # 第一卷沿用原始输出名（保证 output_word 指向有效文件），
                # 第2卷起追加 _partNN 后缀，与第一卷同目录
                if multi_volume and vol_idx > 1:
                    vol_path = f"{out_stem}_part{vol_idx:02d}.docx"
                else:
                    vol_path = output_path
                shutil.move(temp_docx, vol_path)
                output_files.append(vol_path)
                logger.info(
                    f"第{vol_idx}/{len(chunk_starts)}卷完成"
                    f"（第{chunk_start + 1}-{chunk_end}页）: {vol_path} "
                    f"[{os.path.getsize(vol_path) / 1024:.1f}KB]"
                )
            finally:
                # 释放本卷内存与临时目录，再处理下一卷
                chunk_res = None
                last_img = None
                if temp_dir:
                    shutil.rmtree(temp_dir, ignore_errors=True)
                    temp_dir = None
                gc.collect()

        if pages_done:
            log(
                f"OCR 识别完成：{pages_done} 页共 {infer_seconds:.0f}s，"
                f"页均 {infer_seconds / pages_done:.2f}s，总耗时 {time.time() - t_total:.0f}s"
            )

        if not output_files:
            finish(False, "PP-Structure未识别出任何内容")
            return

        if multi_volume:
            logger.info(f"PP-Structure转换完成，共生成{len(output_files)}卷")
            finish(True, (
                f"成功转换（文档较大，共{len(chunk_starts)}卷）：主文件 {output_files[0]}，"
                f"其余分卷（_partNN）在同一目录"
            ))
        else:
            file_size_kb = os.path.getsize(output_files[0]) / 1024
            logger.info(f"PP-Structure转换完成，文件大小: {file_size_kb:.1f}KB")
            finish(True, f"成功转换到: {output_files[0]}")

    except MemoryError:
        finish(False, "内存不足，OCR转换需要较多内存，请尝试转换较少页数")

    except Exception as e:
        logger.error(f"错误详情:\n{traceback.format_exc()}")
        finish(False, f"OCR转换失败: {str(e)}")

    finally:
        # 统一释放资源：异常/正常路径都会执行
        if doc is not None:
            try:
                doc.close()
            except Exception:
                pass
        if temp_dir:
            shutil.rmtree(temp_dir, ignore_errors=True)


class PDFToWordConverter:
    def __init__(self):
        self.converter: Optional["Converter"] = None
        self._is_cancelled = False
        self._process: Optional[multiprocessing.Process] = None

    # ========== 扫描版PDF检测 ==========

    def _is_scanned_pdf(self, pdf_path: str, sample_pages: int = 10) -> bool:
        """
        检测PDF是否为扫描版（每页都是图片，无可提取文字）

        采样前N页，如果所有采样页均无文字且仅有整页图片，则判定为扫描版。

        Args:
            pdf_path: PDF文件路径
            sample_pages: 采样页数

        Returns:
            True 表示是扫描版PDF
        """
        try:
            import fitz
            doc = fitz.open(pdf_path)
            try:
                total = len(doc)
                check_count = min(sample_pages, total)

                if check_count == 0:
                    return False

                no_text_count = 0
                for i in range(check_count):
                    page = doc[i]
                    text = page.get_text().strip()
                    images = page.get_images()
                    if len(text) == 0 and len(images) >= 1:
                        no_text_count += 1
            finally:
                doc.close()

            # 如果所有采样页都无文字且有图片，则判定为扫描版
            is_scanned = (no_text_count == check_count) and (check_count > 0)
            if is_scanned:
                logger.info(f"检测到扫描版PDF（{check_count}/{check_count}页无文字）")
            else:
                logger.info(f"非扫描版PDF（{no_text_count}/{check_count}页无文字）")
            return is_scanned

        except Exception as e:
            logger.info(f"检测PDF类型时出错: {e}")
            return False

    def _convert_scanned_pdf_as_images(
        self,
        pdf_path: str,
        output_path: str,
        start: int = 0,
        end: Optional[int] = None,
        progress_callback=None
    ) -> tuple[bool, str]:
        """
        扫描版PDF的降级转换（图片模式）

        当OCR不可用时，将每页渲染为图片插入Word文档。
        每页使用较高DPI渲染以保证清晰度。

        Args:
            pdf_path: PDF文件路径
            output_path: 输出Word文件路径
            start: 起始页码（从0开始）
            end: 结束页码
            progress_callback: 进度回调函数

        Returns:
            (是否成功, 消息/错误信息)
        """
        import fitz
        import shutil
        from docx import Document
        from docx.shared import Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = None
        temp_dir = None

        try:
            doc = fitz.open(pdf_path)
            total_pages = len(doc)
            actual_end = min(end if end is not None else total_pages, total_pages)
            pages_to_convert = actual_end - start

            logger.info(f"图片模式转换: 共{total_pages}页，转换第{start + 1}-{actual_end}页")

            word_doc = Document()

            temp_dir = tempfile.mkdtemp(prefix='pdf_img_')

            for page_idx in range(start, actual_end):
                if self._is_cancelled:
                    break

                page = doc[page_idx]
                current_page = page_idx - start + 1

                if progress_callback:
                    progress_callback(current_page, pages_to_convert)

                # 渲染页面为图片
                pix = page.get_pixmap(dpi=150)
                img_path = os.path.join(temp_dir, f'page_{page_idx}.png')
                pix.save(img_path)

                if page_idx > start:
                    word_doc.add_page_break()

                # 插入图片到Word
                para = word_doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run()
                run.add_picture(img_path, width=Inches(6.0))

                # 清理临时图片
                try:
                    os.remove(img_path)
                except OSError:
                    pass

                logger.info(f"已处理第 {page_idx + 1}/{total_pages} 页")

            if self._is_cancelled:
                return False, "转换被用户中断"

            word_doc.save(output_path)

            if not os.path.exists(output_path):
                return False, "转换完成但输出文件不存在"

            file_size_kb = os.path.getsize(output_path) / 1024
            logger.info(f"图片模式转换完成，文件大小: {file_size_kb:.1f}KB")
            return True, f"成功转换到: {output_path}"

        except Exception as e:
            error_msg = f"图片模式转换失败: {str(e)}"
            logger.error(f"错误详情:\n{traceback.format_exc()}")
            return False, error_msg

        finally:
            # 统一释放资源：取消/异常/正常路径都会执行
            if doc is not None:
                try:
                    doc.close()
                except Exception:
                    pass
            if temp_dir:
                shutil.rmtree(temp_dir, ignore_errors=True)

    # ========== 扫描版PDF的OCR转换 ==========

    def _convert_scanned_pdf(
        self,
        pdf_path: str,
        output_path: str,
        start: int = 0,
        end: Optional[int] = None,
        progress_callback=None,
        log_callback=None
    ) -> tuple[bool, str]:
        """
        使用 PP-Structure 将扫描版PDF转换为Word文档（独立子进程执行）

        OCR 推理在独立子进程中运行：paddle CPU 推理会吃满所有核心，
        放在 GUI 进程内会拖垮界面；子进程同时设为低优先级，转换期间
        机器保持可用，崩溃也不连累 GUI（与 pdf2docx 隔离策略一致）。

        如果PaddleOCR不可用，回退到图片模式（每页作为图片插入Word）。

        Args:
            pdf_path: PDF文件路径
            output_path: 输出Word文件路径
            start: 起始页码（从0开始）
            end: 结束页码
            progress_callback: 进度回调函数
            log_callback: 可选的界面日志回调（接收一个字符串）

        Returns:
            (是否成功, 消息/错误信息)
        """
        log = log_callback or (lambda m: None)

        # 增强包目录缺失直接走图片模式，避免无谓地拉起子进程
        if _find_ocr_addon_dir() is None:
            log("未找到 OCR 增强包（程序目录下缺少 ocr_addon 文件夹）")
            logger.info("未找到 OCR 增强包，扫描版PDF将使用图片模式转换")
            return self._convert_scanned_pdf_as_images(
                pdf_path, output_path, start, end, progress_callback
            )

        success, msg = self._run_ocr_isolated(
            pdf_path, output_path, start, end, progress_callback, log
        )
        if success:
            return True, msg
        if msg == _OCR_UNAVAILABLE:
            # OCR不可用，回退到图片模式（降级原因已由子进程日志透传）
            log("未启用 OCR，改用图片模式转换扫描版 PDF（输出为图片，无可编辑文字）")
            logger.info("OCR不可用，使用图片模式转换扫描版PDF")
            return self._convert_scanned_pdf_as_images(
                pdf_path, output_path, start, end, progress_callback
            )
        if msg == _OCR_CRASH:
            return False, "OCR 子进程异常退出，无法完成转换"
        logger.info(f"OCR转换失败: {msg}")
        return False, msg

    def _run_ocr_isolated(
        self,
        pdf_path: str,
        output_path: str,
        start: int,
        end: Optional[int],
        progress_callback,
        log
    ) -> tuple[bool, str]:
        """
        在独立子进程中运行 OCR 转换（含 MKLDNN 故障自动降级重试）

        MKLDNN 在部分 Windows CPU 型号上存在原生崩溃的已知案例：默认首次
        尝试开启 MKLDNN，若引擎初始化失败或子进程崩溃，自动降级为关闭
        MKLDNN 重试一次；仍失败则如实返回。设置环境变量
        PDFCONVERTER_NO_MKLDNN=1 可强制全程关闭 MKLDNN。

        Returns:
            (是否成功, 消息；OCR 不可用时为 _OCR_UNAVAILABLE 哨兵)
        """
        if os.environ.get('PDFCONVERTER_NO_MKLDNN', '') == '1':
            attempts = [False]
        else:
            attempts = [True, False]

        result = (False, _OCR_CRASH)
        for idx, use_mkldnn in enumerate(attempts):
            result = self._spawn_ocr_worker(
                pdf_path, output_path, start, end, use_mkldnn,
                progress_callback, log
            )
            if result[0] or idx == len(attempts) - 1:
                return result
            # 仅当失败疑似与 MKLDNN 相关（引擎初始化失败/子进程崩溃）时
            # 才值得降级重试；普通转换失败或用户取消重试无意义
            if result[1] not in (_OCR_UNAVAILABLE, _OCR_CRASH):
                return result
            log("MKLDNN 加速初始化失败或运行中崩溃，自动降级为关闭 MKLDNN 重试...")
        return result

    def _spawn_ocr_worker(
        self,
        pdf_path: str,
        output_path: str,
        start: int,
        end: Optional[int],
        use_mkldnn: bool,
        progress_callback,
        log
    ) -> tuple[bool, str]:
        """
        拉起 OCR 子进程并转发其进度/日志消息，直到子进程结束

        Returns:
            (是否成功, 消息/哨兵)
        """
        ctx = multiprocessing.get_context("spawn")
        msg_queue = ctx.Queue()
        proc = ctx.Process(
            target=_ocr_worker,
            args=(pdf_path, output_path, start, end, msg_queue, use_mkldnn),
        )
        self._process = proc
        proc.start()

        result = None

        def handle(msg):
            nonlocal result
            kind = msg[0]
            if kind == _OCR_MSG_PROGRESS:
                if progress_callback:
                    progress_callback(msg[1], msg[2])
            elif kind == _OCR_MSG_LOG:
                log(msg[1])
            elif kind == _OCR_MSG_RESULT:
                result = (msg[1], msg[2])

        # 轮询排水：转发进度/日志，直到子进程退出且消息排空
        while True:
            try:
                handle(msg_queue.get(timeout=0.3))
            except queue.Empty:
                if not proc.is_alive():
                    break
                continue
        # 子进程退出后再排空残留消息（result 可能晚于进程状态可见）
        while True:
            try:
                handle(msg_queue.get_nowait())
            except queue.Empty:
                break

        proc.join()
        self._process = None

        if self._is_cancelled:
            return False, "转换被用户中断"
        # 子进程异常退出（exitcode 非0/负值表示原生崩溃或被信号终止）
        if proc.exitcode != 0:
            logger.error(f"OCR子进程异常退出: exitcode={proc.exitcode}")
            return False, _OCR_CRASH
        if result is None:
            logger.error("OCR子进程正常退出但未返回结果")
            return False, _OCR_CRASH
        return result

    # ========== 主转换入口 ==========

    def convert(
        self, 
        pdf_path: str, 
        output_path: str,
        start: int = 0,
        end: Optional[int] = None,
        progress_callback=None,
        log_callback=None
    ) -> tuple[bool, str]:
        """
        将PDF转换为Word文档

        自动检测PDF类型：
        - 普通PDF（含可提取文字）：使用pdf2docx转换
        - 扫描版PDF（每页均为图片）：使用OCR识别文字后转换

        Args:
            pdf_path: PDF文件路径
            output_path: 输出Word文件路径
            start: 起始页码（从0开始）
            end: 结束页码（None表示到最后一页）
            progress_callback: 进度回调函数(page_no, total_pages)
            log_callback: 可选的界面日志回调（接收一个字符串），用于透传扫描版检测/OCR状态

        Returns:
            (是否成功, 消息/错误信息)
        """
        log = log_callback or (lambda m: None)
        try:
            # 已收到取消请求则直接退出；不再重置标志，
            # 避免覆盖掉转换启动前到达的取消请求
            # （每个任务使用独立的转换器实例，无需重置）
            if self._is_cancelled:
                return False, "转换已被用户取消"

            if not os.path.exists(pdf_path):
                return False, f"PDF文件不存在: {pdf_path}"

            # 检查文件大小，警告大文件
            file_size_mb = os.path.getsize(pdf_path) / (1024 * 1024)
            if file_size_mb > 50:
                logger.info(f"警告: 文件较大 ({file_size_mb:.1f}MB)，转换可能需要较长时间")

            # 确保输出目录存在
            output_dir = Path(output_path).parent
            output_dir.mkdir(parents=True, exist_ok=True)

            # 检测是否为扫描版PDF
            is_scanned = self._is_scanned_pdf(pdf_path)

            if is_scanned:
                # 扫描版PDF：使用OCR转换
                log("检测到扫描版 PDF（页面无可提取文字层）")
                return self._convert_scanned_pdf(
                    pdf_path, output_path, start, end, progress_callback, log
                )

            # 普通PDF：使用pdf2docx转换（在独立子进程中执行，隔离底层C库崩溃）
            logger.info(f"开始转换: {pdf_path}")
            success, err_msg = self._run_pdf2docx_isolated(
                pdf_path, output_path, start, end
            )
            gc.collect()

            # 转换期间收到取消请求，丢弃结果
            if self._is_cancelled:
                return False, "转换被用户中断"

            if not success:
                return False, err_msg

            # 验证生成的文件
            if not os.path.exists(output_path):
                return False, "转换完成但输出文件不存在"

            if os.path.getsize(output_path) == 0:
                return False, "转换生成的文件为空"

            logger.info(f"转换完成，文件大小: {os.path.getsize(output_path) / 1024:.1f}KB")

            # 后处理：修复图片显示问题
            try:
                self._fix_images_in_docx(output_path)
            except Exception as e:
                logger.info(f"图片修复失败（不影响主文件）: {e}")

            return True, f"成功转换到: {output_path}"

        except MemoryError:
            if self.converter:
                try:
                    self.converter.close()
                except Exception:
                    pass
            self.converter = None
            gc.collect()
            return False, "内存不足，请尝试转换较小的PDF文件或关闭其他程序"

        except KeyboardInterrupt:
            if self.converter:
                try:
                    self.converter.close()
                except Exception:
                    pass
            self.converter = None
            return False, "转换被用户中断"

        except Exception as e:
            error_msg = f"转换失败: {str(e)}"
            logger.error(f"错误详情:\n{traceback.format_exc()}")

            if self.converter:
                try:
                    self.converter.close()
                except Exception:
                    pass
            self.converter = None
            gc.collect()

            return False, error_msg
    
    def _fix_images_in_docx(self, docx_path: str):
        """
        修复Word文档中的图片显示问题
        
        问题1：图片显示不全或显示为空白
        问题2：某些白色/透明图片遮挡了其他内容
        问题3：水印、分隔线等装饰性元素（几乎纯白色、内容极少）
        问题4：anchor类型浮动图片尺寸异常（如高度数千英寸），导致Word无法打开
        
        解决方案：
        1. 检测并删除异常尺寸的浮动图片（anchor类型）
        2. 检测并删除几乎纯白色的图片（水印/分隔线）
        3. 删除异常大的图片
        4. 调整正常图片的尺寸
        """
        try:
            from docx import Document
            from docx.shared import Inches, Emu
            from docx.oxml.ns import qn
            from PIL import Image
            import io
            
            doc = Document(docx_path)
            body = doc.element.body
            
            # ========== 第一步：处理anchor类型的浮动图片 ==========
            # anchor图片可能是pdf2docx生成的背景层/水印
            # 这些图片如果尺寸异常（如高度数千英寸），会导致Word无法打开
            self._fix_anchor_drawings(body, doc)
            
            # ========== 第二步：处理inline类型的图片 ==========
            # 直接收集待删图片的XML节点引用，避免与anchor图片的索引空间混淆
            shapes_to_remove = []
            
            for idx, shape in enumerate(doc.inline_shapes):
                try:
                    width_px = shape.width
                    height_px = shape.height
                    
                    is_blank_image = False
                    
                    try:
                        image_part = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
                        image_part_rid = image_part
                        image_part = doc.part.related_parts.get(image_part_rid)
                        
                        if image_part:
                            img = Image.open(io.BytesIO(image_part.blob))
                            img_rgb = img.convert('RGB')

                            # 缩略图分析：缩小到100x100后遍历，避免逐像素getpixel
                            small_img = img_rgb.resize((100, 100), Image.LANCZOS)
                            samples = list(small_img.getdata())
                            small_img.close()

                            if samples:
                                avg_r = sum(p[0] for p in samples) / len(samples)
                                avg_g = sum(p[1] for p in samples) / len(samples)
                                avg_b = sum(p[2] for p in samples) / len(samples)
                                
                                min_channel = min(avg_r, avg_g, avg_b)
                                max_channel = max(avg_r, avg_g, avg_b)
                                
                                if min_channel > 240 and (max_channel - min_channel) < 15:
                                    is_blank_image = True
                                    logger.info(f"图片 {idx}: 检测到几乎纯白色图片（水印/分隔线），已删除")
                            
                            img.close()
                            img_rgb.close()
                        
                    except Exception as e:
                        logger.info(f"分析图片 {idx} 时出错: {e}")
                    
                    if is_blank_image:
                        shapes_to_remove.append(shape._inline)
                        continue
                    
                    max_reasonable_size = Inches(8)
                    if width_px > max_reasonable_size or height_px > max_reasonable_size:
                        shapes_to_remove.append(shape._inline)
                        logger.info(f"图片 {idx}: 尺寸异常大，已删除")
                        continue
                    
                    if width_px > 0 and height_px > 0:
                        aspect_ratio = width_px / height_px
                        if aspect_ratio > 10 or aspect_ratio < 0.1:
                            shapes_to_remove.append(shape._inline)
                            logger.info(f"图片 {idx}: 宽高比异常（{aspect_ratio:.2f}），可能是分隔线，已删除")
                            continue
                    
                    max_width = Inches(6.5)
                    if shape.width > max_width:
                        aspect_ratio = shape.height / shape.width
                        shape.width = max_width
                        shape.height = max_width * aspect_ratio
                        
                except Exception as e:
                    logger.info(f"处理图片 {idx} 时出错: {e}")
                    continue
            
            if shapes_to_remove:
                logger.info(f"发现 {len(shapes_to_remove)} 个问题inline图片，正在清理...")
                self._remove_problematic_shapes(shapes_to_remove)
            
            doc.save(docx_path)
            
        except Exception as e:
            logger.info(f"图片修复警告: {e}")
    
    def _fix_anchor_drawings(self, body, doc=None):
        """
        处理anchor类型的浮动图片

        pdf2docx 会将PDF中的背景图层转为anchor类型的浮动图片（behindDoc=1），
        这些背景层里往往包含页面上的图表/截图，不能误删。

        处理策略：
        1. behindDoc=1 且有实质内容 → 保留（保持 behindDoc=1，作为正文背景）
        2. behindDoc=1 且为完全透明的占位图 → 删除
        3. 尺寸异常的巨型图片 → 删除
        """
        try:
            from docx.shared import Inches
            from docx.oxml.ns import qn

            ns_wp = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
            ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
            ns_r = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

            # 异常巨型图片阈值：pdf2docx 会把整页背景渲染为接近页面尺寸的
            # behindDoc 图片，阈值过小会把这些正常背景图误删，所以放宽到
            # 明显超出常规页面很多的尺寸（Word 仍可正常打开）。
            max_reasonable_width_emu = int(Inches(40))
            max_reasonable_height_emu = int(Inches(60))

            # 页面尺寸，用于把页背景图缩放到 Word 页面大小
            try:
                page_width_emu = int(doc.sections[0].page_width)
                page_height_emu = int(doc.sections[0].page_height)
            except Exception:
                page_width_emu = int(Inches(8.27))
                page_height_emu = int(Inches(11.69))

            # 判断页面背景图的尺寸下限（大于此尺寸才认为是整页背景）
            min_page_width_emu = int(Inches(10))
            min_page_height_emu = int(Inches(10))

            # 查找所有 drawing 元素（使用 .// 递归搜索嵌套元素）
            drawings = body.findall('.//' + qn('w:drawing'))
            removed_count = 0
            behind_removed_count = 0
            behind_preserved_count = 0

            for drawing in drawings:
                try:
                    # 检查是否有anchor子元素
                    anchor = drawing.find(qn('wp:anchor'))
                    if anchor is None:
                        anchor = drawing.find('{' + ns_wp + '}anchor')

                    if anchor is None:
                        continue

                    # 检查是否是背景层（behindDoc=true）
                    behind_doc = anchor.get('behindDoc', 'false')
                    is_behind_doc = behind_doc == '1' or behind_doc.lower() == 'true'

                    # 获取anchor的extent（尺寸）
                    extent = anchor.find(qn('wp:extent'))
                    if extent is None:
                        extent = anchor.find('{' + ns_wp + '}extent')

                    if extent is not None:
                        cx = int(extent.get('cx', '0'))
                        cy = int(extent.get('cy', '0'))
                        is_huge = cx > max_reasonable_width_emu or cy > max_reasonable_height_emu
                    else:
                        cx = 0
                        cy = 0
                        is_huge = False

                    should_remove = False

                    if is_huge:
                        # 异常巨型浮动图片，无论是否behindDoc都删除
                        cx_inches = cx / 914400
                        cy_inches = cy / 914400
                        logger.info(f"删除异常巨型浮动图片: {cx_inches:.1f}x{cy_inches:.1f} 英寸, behindDoc={behind_doc}")
                        should_remove = True
                    elif is_behind_doc:
                        # behindDoc=1 的图片通常是 pdf2docx 生成的页面背景层，
                        # 里面往往包含正文图表，不能按水印逻辑误删。
                        # 只删除完全透明或极小的占位图，其余保留在文字下方。
                        is_empty = self._is_transparent_image(drawing, doc, ns_a, ns_r)
                        if is_empty:
                            width_inches = cx / 914400
                            height_inches = cy / 914400
                            logger.info(f"删除透明/占位背景图片: {width_inches:.1f}x{height_inches:.1f} 英寸")
                            should_remove = True
                            behind_removed_count += 1
                        else:
                            # 有实质内容的 behindDoc 图片作为正文背景保留，
                            # 保持 behindDoc=1 以免遮挡已提取的可编辑文字。
                            # 若它是接近页面尺寸的整页背景，则缩放到 Word 页面大小，
                            # 避免原图 20+ 英寸导致显示不全/只显示左上角。
                            if (cx >= min_page_width_emu and cy >= min_page_height_emu):
                                self._scale_anchor_to_page(
                                    anchor, page_width_emu, page_height_emu, ns_wp
                                )
                            behind_preserved_count += 1

                    if should_remove:
                        parent = drawing.getparent()
                        if parent is not None:
                            parent.remove(drawing)
                            removed_count += 1

                except Exception as e:
                    logger.info(f"处理anchor图片时出错: {e}")
                    continue

            if removed_count > 0:
                logger.info(f"共删除 {removed_count} 个浮动图片（其中背景水印 {behind_removed_count} 个）")
            if behind_preserved_count > 0:
                logger.info(f"保留 {behind_preserved_count} 个含内容的背景浮动图片")

        except Exception as e:
            logger.info(f"修复浮动图片时出错: {e}")

    def _scale_anchor_to_page(self, anchor, page_width_emu, page_height_emu, ns_wp):
        """
        把整页背景 anchor 缩放到 Word 页面大小，并重置定位到页面左上角

        pdf2docx 生成的页背景图通常被渲染成 20+ 英寸的巨型 anchor，并带
        有负的 positionOffset，Word 默认只显示左上角一部分。将其缩放到真实
        页面尺寸并定位到 (0,0) 可保证整页内容可见且比例正确。
        """
        try:
            from docx.oxml.ns import qn

            # 保持 behindDoc=1，确保图片在文字下方
            anchor.set('behindDoc', '1')

            # 修改 anchor 外层 extent
            extent = anchor.find(qn('wp:extent'))
            if extent is None:
                extent = anchor.find('{' + ns_wp + '}extent')
            if extent is not None:
                extent.set('cx', str(page_width_emu))
                extent.set('cy', str(page_height_emu))

            # 重置 position 偏移为 0
            for pos_tag in ('positionH', 'positionV'):
                pos = anchor.find(qn('wp:' + pos_tag))
                if pos is None:
                    pos = anchor.find('{' + ns_wp + '}' + pos_tag)
                if pos is not None:
                    off = pos.find(qn('wp:posOffset'))
                    if off is None:
                        off = pos.find('{' + ns_wp + '}posOffset')
                    if off is not None:
                        off.text = '0'

            # simplePos 也归零
            simple_pos = anchor.find(qn('wp:simplePos'))
            if simple_pos is None:
                simple_pos = anchor.find('{' + ns_wp + '}simplePos')
            if simple_pos is not None:
                simple_pos.set('x', '0')
                simple_pos.set('y', '0')

            # 同步修改内层图形的 extent/off，避免内层尺寸与外层不匹配导致裁剪
            ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
            for ext in anchor.findall('.//' + qn('a:ext')):
                ext.set('cx', str(page_width_emu))
                ext.set('cy', str(page_height_emu))
            for off in anchor.findall('.//' + qn('a:off')):
                off.set('x', '0')
                off.set('y', '0')

        except Exception as e:
            logger.info(f"缩放页背景图时出错: {e}")

    def _is_transparent_image(self, drawing, doc, ns_a, ns_r):
        """
        判断 behindDoc 背景图片是否为完全透明的占位图

        pdf2docx 在某些文档里会生成大量全透明的背景层（例如 RGBA
        图片但 alpha 通道全为 0）。这类图片没有可见内容，可以安全删除，
        避免在 Word 中产生冗余的浮动对象。

        Args:
            drawing: XML drawing 元素
            doc: Document 对象，用于获取图片数据
            ns_a: drawingml 主命名空间
            ns_r: 关系命名空间

        Returns:
            True 表示是完全透明的占位图，应删除；False 保留
        """
        try:
            from docx.oxml.ns import qn
            from PIL import Image
            import io

            # 查找图片引用
            blip = drawing.find('.//' + qn('a:blip'))
            if blip is None:
                blip = drawing.find('.//{' + ns_a + '}blip')

            if blip is None:
                return True

            embed = blip.get(qn('r:embed'))
            if embed is None:
                embed = blip.get('{' + ns_r + '}embed')

            if not embed or doc is None:
                return False

            image_part = doc.part.related_parts.get(embed)
            if image_part is None:
                return False

            blob_size = len(image_part.blob)
            if blob_size < 1000:
                return True

            img = Image.open(io.BytesIO(image_part.blob))
            if img.mode == 'RGBA':
                alpha = img.split()[-1]
                # 如果 alpha 最大值极低，认为完全透明
                if alpha.getextrema()[1] < 10:
                    img.close()
                    alpha.close()
                    return True
                alpha.close()

            img.close()
            return False

        except Exception:
            return False
    
    def _remove_problematic_shapes(self, inline_elements):
        """
        从文档中删除指定的inline图片
        
        直接基于inline shape的XML节点删除其所在的w:drawing元素，
        避免使用findall索引（会与anchor图片混在同一列表导致错位误删）
        """
        try:
            from docx.oxml.ns import qn
            
            for inline_elem in inline_elements:
                # inline 元素的父节点即 w:drawing
                drawing = inline_elem.getparent()
                if drawing is None or drawing.tag != qn('w:drawing'):
                    continue
                parent = drawing.getparent()
                if parent is not None:
                    parent.remove(drawing)
                        
        except Exception as e:
            logger.info(f"删除问题图片时出错: {e}")
    
    def _run_pdf2docx_isolated(self, pdf_path, output_path, start, end):
        """
        在独立子进程中运行 pdf2docx 转换，隔离底层C库崩溃。

        Returns:
            (是否成功, 错误信息)
        """
        ctx = multiprocessing.get_context("spawn")
        result_queue = ctx.Queue()
        proc = ctx.Process(
            target=_pdf2docx_worker,
            args=(pdf_path, output_path, start, end, result_queue),
        )
        self._process = proc
        proc.start()
        proc.join()
        self._process = None

        # 被取消（子进程已被 terminate）
        if self._is_cancelled:
            return False, "转换被用户中断"

        # 子进程异常退出（exitcode 非0/负值表示原生崩溃或被信号终止）
        if proc.exitcode != 0:
            logger.error(f"pdf2docx子进程异常退出: exitcode={proc.exitcode}")
            return False, "该PDF触发底层解析库崩溃，无法转换（文档结构复杂或不兼容）"

        # 读取子进程返回的结果
        try:
            return result_queue.get_nowait()
        except queue.Empty:
            # 队列为空但进程正常退出，回退到检查输出文件
            return os.path.exists(output_path), "转换结果未知"

    def cancel(self):
        """
        取消转换（协作式）

        设置取消标志；若 pdf2docx 子进程正在运行则直接终止它。
        禁止跨线程关闭主进程内的 C 层对象，避免数据竞争导致崩溃。
        """
        self._is_cancelled = True
        proc = self._process
        if proc is not None and proc.is_alive():
            try:
                proc.terminate()
            except Exception:
                pass
